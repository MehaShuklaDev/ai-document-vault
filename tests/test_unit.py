"""Pure unit tests — no database, no network."""

from __future__ import annotations

import uuid

import pytest

from app.ai import rag
from app.ai.analysis import DocumentAnalysis, SummaryOptions, build_excerpt
from app.ai.prompts import analysis_messages, chat_messages, format_context
from app.ai.providers import FakeEmbeddings, FakeLLM, Message, estimate_cost, parse_json_response
from app.ai.retrieval import RetrievedChunk, rrf_fuse, trim_to_budget
from app.services.chunking import chunk_pages, count_tokens
from app.services.extraction import MIME_DOCX, MIME_MD, MIME_PDF, MIME_TXT, Page, clean_text, extract_pages, sniff_mime

# --------------------------------------------------------------------- chunking


def test_chunker_respects_budget_and_overlap():
    text = " ".join(f"Sentence number {i} talks about topic {i % 7}." for i in range(400))
    chunks = chunk_pages([Page(1, text)], chunk_tokens=120, overlap_tokens=20)
    assert len(chunks) > 3
    for c in chunks:
        assert c.token_count <= 120 + 25  # overlap + join slack
    # overlap: the tail of chunk n appears at the head of chunk n+1
    tail = chunks[0].content[-40:]
    assert tail.split()[-1] in chunks[1].content[:200]
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_chunker_tracks_pages_and_sections():
    pages = [Page(1, "# Intro\n\nFirst paragraph on page one."), Page(2, "Second page paragraph, still intro."), Page(3, "RISKS\n\nA risk paragraph.")]
    chunks = chunk_pages(pages, chunk_tokens=400, overlap_tokens=20)
    assert chunks[0].page_start == 1 and chunks[0].page_end == 3
    assert chunks[0].section == "Intro"


def test_chunker_splits_giant_paragraph():
    text = "word " * 3000  # one paragraph, no sentence boundaries
    chunks = chunk_pages([Page(1, text)], chunk_tokens=200, overlap_tokens=10)
    assert all(c.token_count <= 215 for c in chunks)
    assert len(chunks) >= 10


def test_chunker_max_chunks_guard():
    text = "\n\n".join(f"Paragraph {i} with some words in it." for i in range(500))
    chunks = chunk_pages([Page(1, text)], chunk_tokens=60, overlap_tokens=5, max_chunks=7)
    assert len(chunks) == 7


def test_chunker_rejects_bad_overlap():
    with pytest.raises(ValueError):
        chunk_pages([Page(1, "x")], chunk_tokens=10, overlap_tokens=10)


def test_count_tokens_monotonic():
    assert count_tokens("hello world") < count_tokens("hello world, and a lot more words here")


# ------------------------------------------------------------------- extraction


def test_sniff_mime_prefers_magic_bytes():
    assert sniff_mime(b"%PDF-1.7 ...", "x.txt", "text/plain") == MIME_PDF
    assert sniff_mime(b"PK\x03\x04....", "report.docx", None) == MIME_DOCX
    assert sniff_mime(b"# Title\n", "notes.md", None) == MIME_MD
    assert sniff_mime(b"plain words", "notes.txt", None) == MIME_TXT
    assert sniff_mime(b"\x00\x01\x02binary", "blob.bin", None) == "application/octet-stream"


def test_extract_text_paginates_long_files():
    data = ("lorem ipsum " * 2000).encode()
    pages = extract_pages(data, MIME_TXT)
    assert len(pages) > 1 and pages[0].number == 1 and pages[-1].number == len(pages)


def test_clean_text_normalises_whitespace():
    assert clean_text("a \t b\r\n\r\n\r\n\r\nc  ") == "a b\n\nc"


def test_extract_pdf_roundtrip():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    import io

    buf = io.BytesIO()
    w.write(buf)
    # blank page → no text → ExtractionError (scanned-PDF path)
    from app.services.extraction import ExtractionError

    with pytest.raises(ExtractionError):
        extract_pages(buf.getvalue(), MIME_PDF)


def test_extract_docx_headings_become_sections():
    import io

    import docx

    d = docx.Document()
    d.add_heading("Vacation Policy", level=1)
    d.add_paragraph("Employees accrue 1.5 days per month.")
    d.add_heading("Remote Work", level=1)
    d.add_paragraph("Three days per week.")
    buf = io.BytesIO()
    d.save(buf)
    pages = extract_pages(buf.getvalue(), MIME_DOCX)
    assert [p.section for p in pages] == ["Vacation Policy", "Remote Work"]


# -------------------------------------------------------------------- retrieval


def _chunk(i: int, tokens: int = 100) -> RetrievedChunk:
    return RetrievedChunk(
        id=uuid.uuid4(),
        document_id=uuid.uuid4(),
        version_id=uuid.uuid4(),
        chunk_index=i,
        content=f"c{i}",
        token_count=tokens,
        page_start=1,
        page_end=1,
        section=None,
        score=0,
    )


def test_rrf_rewards_agreement():
    a, b, c = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    scores = rrf_fuse([[a, b, c], [b, c, a]], k=60)
    assert scores[b] > scores[a] and scores[b] > scores[c]  # b is ranked 2 and 1 → best combined
    assert scores[a] == pytest.approx(1 / 61 + 1 / 63)


def test_trim_to_budget_keeps_at_least_one():
    chunks = [_chunk(0, 500), _chunk(1, 100)]
    assert len(trim_to_budget(chunks, 200)) == 1
    assert len(trim_to_budget([_chunk(0, 100), _chunk(1, 100), _chunk(2, 100)], 250)) == 2


# ---------------------------------------------------------------------- RAG glue


def test_verify_citations_drops_hallucinated_and_renumbers():
    chunks = [_chunk(0), _chunk(1), _chunk(2)]
    names = {c.document_id: f"doc{i}.pdf" for i, c in enumerate(chunks)}
    cits = rag.build_citations(chunks, names)
    text, kept = rag.verify_citations("Revenue rose 18% [3]. Margin improved [7]. Dividend approved [3][1].", cits)
    assert text == "Revenue rose 18% [1]. Margin improved. Dividend approved [1][2]."
    assert [c.n for c in kept] == [1, 2]
    assert kept[0].chunk_id == chunks[2].id and kept[1].chunk_id == chunks[0].id


def test_format_context_and_chat_prompt_shape():
    ctx = format_context([(1, "a.pdf, p.2", "Alpha text"), (2, "b.docx", "Beta text")])
    msgs = chat_messages(ctx, [("user", "hi"), ("assistant", "hello")], "What is alpha?")
    assert msgs[0].role == "system" and "[1] (a.pdf, p.2)" in msgs[0].content
    assert [m.role for m in msgs] == ["system", "user", "assistant", "user"]
    assert "ONLY the numbered context" in msgs[0].content


def test_analysis_prompt_interpolates_options_without_breaking_json_braces():
    m = analysis_messages("body", length="bullets", tone="executive", focus="risks")
    assert "5-8 bullet points" in m[0].content and "focus especially on: risks" in m[0].content
    assert '"key_points"' in m[0].content  # literal braces survived str.format


# -------------------------------------------------------------------- analysis


def test_build_excerpt_head_plus_samples_within_budget():
    chunks = [f"chunk {i}" for i in range(50)]
    tokens = [100] * 50
    ex = build_excerpt(chunks, tokens, budget=1000)
    assert ex.startswith("chunk 0") and "chunk 1" in ex and "[...]" in ex
    assert sum(1 for i in range(50) if f"chunk {i}\n" in ex + "\n") <= 10


def test_document_analysis_normalises():
    a = DocumentAnalysis.model_validate(
        {"title": "T", "summary": "S", "category": "Weird", "tags": ["A", "a", "b"] * 5, "sentiment": {"label": "positive", "score": 0.4, "rationale": ""}}
    )
    assert a.category == "other" and a.tags == ["a", "b"]


def test_summary_options_hash_is_stable():
    assert SummaryOptions(length="short", tone="neutral").hash() == SummaryOptions(tone="neutral", length="short").hash()
    assert SummaryOptions(length="short").hash() != SummaryOptions(length="long").hash()


# -------------------------------------------------------------------- providers


def test_fake_embeddings_similarity_is_meaningful():
    emb = FakeEmbeddings(256)
    v, usage = emb.embed(["quarterly revenue grew", "revenue growth this quarter", "vacation policy for employees"])
    dot = lambda a, b: sum(x * y for x, y in zip(a, b, strict=True))
    assert dot(v[0], v[1]) > dot(v[0], v[2])
    assert usage.provider == "fake" and len(v[0]) == 256


def test_fake_llm_grounds_answer_in_context():
    ctx = format_context(
        [(1, "r.pdf, p.1", "Revenue for the third quarter reached $12.4 million, up 18%."), (2, "r.pdf, p.2", "The Berlin office is led by Maria Fontaine.")]
    )
    msgs = chat_messages(ctx, [], "Who leads the Berlin office?")
    out = FakeLLM().complete(msgs).text
    assert "Maria Fontaine" in out and "[2]" in out


def test_fake_llm_json_mode_matches_analysis_schema():
    text = FakeLLM().complete(analysis_messages("ACME report body"), json_mode=True).text
    DocumentAnalysis.model_validate(parse_json_response(text))


def test_parse_json_tolerates_fences_and_prose():
    assert parse_json_response('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json_response('Sure! Here it is: {"a": [1,2]} hope that helps') == {"a": [1, 2]}


def test_cost_estimate():
    assert estimate_cost("gpt-4o-mini", 1_000_000, 0) == pytest.approx(0.15)
    assert estimate_cost("text-embedding-3-small", 500_000, 0) == pytest.approx(0.01)
    assert estimate_cost("fake-llm", 10, 10) == 0
    assert estimate_cost("unknown-model", 1_000_000, 0) > 0  # conservative fallback


@pytest.mark.asyncio
async def test_fake_llm_stream_yields_usage_last():
    items = [
        i
        async for i in FakeLLM().astream(
            [Message("system", "CONTEXT PASSAGES\n[1] (x)\nSome passage text here that is long enough."), Message("user", "passage?")]
        )
    ]
    assert isinstance(items[-1].__class__.__name__, str) and items[-1].__class__.__name__ == "Usage"
    assert "".join(i for i in items[:-1] if isinstance(i, str)).strip()


def test_claude_cli_provider_parses_envelope(monkeypatch):
    import subprocess

    from app.ai.providers import ClaudeCLILLM

    captured = {}

    def fake_run(cmd, **kw):
        captured["cmd"] = cmd
        payload = {
            "result": '{"ok": true}',
            "is_error": False,
            "total_cost_usd": 0.0015,
            "usage": {"input_tokens": 1400, "cache_creation_input_tokens": 70, "cache_read_input_tokens": 3, "output_tokens": 9},
            "modelUsage": {"claude-haiku-4-5": {}},
        }
        import json as _json

        return subprocess.CompletedProcess(cmd, 0, stdout=_json.dumps(payload), stderr="")

    monkeypatch.setattr(subprocess, "run", fake_run)
    llm = ClaudeCLILLM("haiku", "/nonexistent/claude", 30)
    c = llm.complete([Message("system", "Sys."), Message("user", "hi"), Message("assistant", "yo"), Message("user", "again")], json_mode=True)
    assert c.text == '{"ok": true}'
    assert c.usage.model == "claude-cli:claude-haiku-4-5" and c.usage.input_tokens == 1473 and c.usage.output_tokens == 9
    assert c.usage.cost_usd == 0.0015  # provider-reported cost wins over the price table
    cmd = captured["cmd"]
    assert "--tools" in cmd and cmd[cmd.index("--tools") + 1] == "" and "--output-format" in cmd
    assert "JSON object" in cmd[cmd.index("--system-prompt") + 1]
    assert "USER: again" in cmd[cmd.index("-p") + 1]


def test_claude_cli_provider_surfaces_failures(monkeypatch):
    import subprocess

    from app.ai.providers import ClaudeCLILLM, LLMError

    monkeypatch.setattr(subprocess, "run", lambda cmd, **kw: subprocess.CompletedProcess(cmd, 1, stdout="", stderr="not logged in"))
    with pytest.raises(LLMError, match="not logged in"):
        ClaudeCLILLM("haiku", "/nonexistent/claude", 30).complete([Message("user", "hi")])


def test_chat_history_strips_stale_citation_markers():
    msgs = chat_messages("ctx", [("user", "q1 [1]?"), ("assistant", "Revenue was $12.4M [1]. Berlin lead is Maria [2][3].")], "q2")
    assert msgs[2].content == "Revenue was $12.4M. Berlin lead is Maria."
    assert msgs[1].content == "q1 [1]?"  # user text untouched
    assert "Earlier turns used their own numbering" in msgs[0].content


def test_provider_auto_resolution(monkeypatch):
    from app.core.config import Settings

    base = dict(llm_provider="auto", embedding_provider="auto", openai_api_key=None, anthropic_api_key=None, _env_file=None)
    s = Settings(**base)
    monkeypatch.setattr(Settings, "find_claude_cli", lambda self: None)
    from app.core.config import ProviderNotConfigured

    with pytest.raises(ProviderNotConfigured, match="OPENAI_API_KEY or ANTHROPIC_API_KEY"):
        _ = s.resolved_llm_provider
    assert s.resolved_embedding_provider == "fake"  # local hashed embeddings are fine without a key
    monkeypatch.setattr(Settings, "find_claude_cli", lambda self: "/usr/local/bin/claude")
    assert s.resolved_llm_provider == "claude-cli"
    assert Settings(**{**base, "anthropic_api_key": "k"}).resolved_llm_provider == "anthropic"
    s2 = Settings(**{**base, "openai_api_key": "k", "anthropic_api_key": "k"})
    assert (s2.resolved_llm_provider, s2.resolved_embedding_provider) == ("openai", "openai")
    assert Settings(**{**base, "llm_provider": "fake", "openai_api_key": "k"}).resolved_llm_provider == "fake"  # explicit wins
